# -*- coding: utf-8 -*-
"""验证板任务分解 (4 级) → 验证板任务分解_2026-10-10.xlsx
改进度就改下面 ROWS 里的「状态/完成度/证据/下一步」, 然后:  cmd.exe /c "python 验证板任务分解_mkxlsx.py"
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '验证板任务分解_2026-10-10.xlsx')

# ---------------- 一级目标参数 ----------------
PARAMS = [
    ('屏', '160×180 px, 点距 0.9375 mm (150×168.75 mm), ICND2260, 60 扫', 'SCH-0.9375-2260-60S-LVDS'),
    ('切片 (名义, 已不用)', '503 片/圈 = 2π×80 px', 'project_pov3d_trial_board_503_slices_60scan_ok'),
    ('转速 / 体帧率 (名义)', '1050 RPM = 17.5 rev/s ⇒ 体帧率 35 Hz (= 2×转速)', 'project_pov3d_final_target 09-22 口径'),
    ('名义需求', '503 × 17.5 = 8,802 slice/s/面', ''),
    ('🔴 硬约束', '60 扫屏只能 1串2 (1串1 已实测作废) ⇒ 每面上限 ≈4,300 fps @DCLK 166.67 MHz', 'project_icnd2260_pairing_mode_bench_0921'),
    ('✅ 定案 (10-10 用户)', '体帧率 24 Hz = 12 rev/s = 720 RPM, 358 片/圈 ⇒ 4,296 slice/s/面 (跑满 4,300 上限, 帧周期 232.8 µs)\n'
                         '角分辨率 ≈1.0°; 最外圈弧向间距 ≈1.40 px (1.31 mm), 半径 ≈57 px 处与径向持平\n'
                         '⚠ 4,300 是 DCLK 166.67 MHz 的数; 桥若只到 150 MHz 上限 ≈3,870 ⇒ 24 Hz 只能 ≈322 片', '用户 10-10 定 (原候选 A/B/C 作废)'),
    ('色深 / 链路', 'RGB565 16 bit/px; 上行 8 bit/通道, 桥片 8→16 展开', ''),
    ('DDR 需求 (双面)', '57,600 B/片 × 4,296 × 2 ≈ 0.50 GB/s;  DR1M90 实测上限 1.79 GB/s ⇒ 余量足', 'X1 2026-10-10'),
    ('主控 / 桥片', 'DR1M90 (Cortex-A35 双核 + PL) ×1, GW5A-25 ×2 (每面一颗)', ''),
    ('跨转子', '无线供电 48 W + WiFi + 光电索引 (本表不含硬件)', ''),
]

# ---------------- 任务树 ----------------
# (编号, 2级, 3级, 4级, 内容/判据, 状态, 完成度%, 验证层级, 依赖硬件, 脱离硬件?, 独立验证?, 证据/记录, 下一步)
S_OK, S_DO, S_NO, S_BLK, S_SKIP = '✅完成', '🔶进行中', '⬜未开始', '⛔阻塞', '⏸跳过'
ROWS = [
 # ===== 1.1 DR1M90 PS 软件 =====
 ('1.1', 'DR1M90 · PS 软件', '', '', 'Linux 侧: 加载 PL / 帧区驱动 / 收流解码 / 控制 / 开机 / 远程', S_DO, 45, '', '', '', '', '', ''),
 ('1.1.1', '', '启动底座', '', 'aarch64 buildroot、configfs 挂载、fpga_manager 加载 PL、GP0 不挂死', S_OK, 100, '自测', 'DR1M90', '否', '是', 'project_dr1m90_gp0_first_light; D2 LOAD_POV_M90 PASS (10-10)', ''),
 ('1.1.2', '', 'povmem 帧区驱动', '', 'arm64 写合并映射; 判据 cpy ≤26 ms', S_OK, 100, '自测', 'DR1M90', '否', '是', 'D3 cpy 11.0 ms (10-10), dr1v90 1874fc2', ''),
 ('1.1.3', '', 'povctl 控制工具', '', 'aarch64 版; 寄存器/fb/fake/oeprobe', S_DO, 90, '自测', 'DR1M90 (+2047 屏看图)', '否', '是', '无屏空跑: OE 翻转, duty 0.041 (10-10)', 'D4 接屏看图'),
 ('1.1.4', '', 'pov_rxd 收流/解码/翻页', '', '本机 lz4 帧能翻页 → WiFi 收流; 连推 ≥200 帧不卡', S_NO, 30, '自测→合测', 'DR1M90 (+屏)', '编译可脱离', '是', '能编不能跑 (缺 povmem+显示 PL, 现已具备)', 'D6'),
 ('1.1.5', '', 'WiFi 推流链路', '', 'PVS1 / povsh; 静止台架 flip 稳定', S_DO, 50, '合测', 'DR1M90 + AP + PC', '否', '否', 'WiFi TCP 159 Mbps (project_dr1m90_wifi_retest)', 'D7'),
 ('1.1.6', '', '开机自启 / 打包', '', 'povboot AUTOBOOT: 冷启动一次到可推流; 打包脚本按架构选', S_DO, 60, '自测', 'DR1M90', '脚本已离线完成', '是', 'tools/pov_arch.sh + 6 个脚本 aarch64 化, dry-run 两架构 PASS; povsh_aarch64 已出', 'D18: aarch64 rootfs (S95pov) + povboot 加载后写 0x33'),
 ('1.1.7', '', '远程状态回传', '', '只连 WiFi 看 fps/miss/桥状态/误码', S_NO, 0, '合测', 'DR1M90 + WiFi', '协议设计可脱离', '是', '', 'D19'),
 ('1.1.8', '', 'JTAG 给 GW5A 灌 SRAM', '', 'bridge_load aarch64; 判据断言 ARM aarch64', S_NO, 10, '合测', 'DR1M90 + ACG525 + 飞线', '编译可脱离', '否', 'project_dr1m90_jtag_port_checklist', 'D14'),
 ('1.1.9', '', 'DDR 带宽 (X1)', '', '双口切片读', S_OK, 100, '自测', 'DR1M90', '否', '是', '1786 MB/s η0.837; 验证板需 ≈0.5 GB/s 余量足', ''),
 # ===== 1.2 DR1M90 PL =====
 ('1.1.10', '', '板端体素→切片', '', '切片放 A35 还是 PL: CPU 路径 DDR 占 83% (12 fps) / PL 直供 28%; 列中心半像素、面 B 相位只加一次、gamma/径向补偿', S_NO, 0, '自测', 'DR1M90', '算法已离线对拍', '是', 'agent C 10-10: vb_slice_all.c x86/aarch64 可编, 8 条板端待决问题见 project_host_validation_profile_358', 'D6/D16; 建议 PL 直供'),
 ('1.1.11', '', '远程升级 (X4)', '', 'SD 卡 A/B 槽位; 热切换 + 自检窗口 + 自动回滚; 开机 trial 回退; GW5A 每次上电由 DR1 灌, 不写 Flash; 内核第一期不做', S_NO, 0, '自测→系统', 'DR1M90 + WiFi (+ACG525 验 GW5A 灌)', '工具/脚本可脱离', '是', '计划 §10 (10-10 用户定)', 'D18–D19 之后, ≈2.5 天'),
 ('1.1.11.a', '', '', '槽位 + MANIFEST + PC 推送工具', 'povota push 只写非当前槽, sha256 校验', S_NO, 0, '自测', 'DR1M90', '大部分可', '是', '', '0.5 天'),
 ('1.1.11.b', '', '', '热切换 + 自检 + 自动回滚', 'GP0 MAGIC / link_up / 好帧计数 / PC 心跳', S_NO, 0, '合测', 'DR1M90 + WiFi', '否', '是', '', '1 天'),
 ('1.1.11.c', '', '', '开机 trial 回退', 'trial 未清 ⇒ 起旧槽位', S_NO, 0, '自测', 'DR1M90', '否', '是', '', '0.5 天'),
 ('1.1.11.d', '', '', '故障注入', '坏包 / 自检失败 / 切换中掉电 各 ≥5 次不变砖', S_NO, 0, '系统', 'DR1M90 + WiFi', '否', '是', '', '0.5 天'),
 ('1.2', 'DR1M90 · PL 逻辑', '', '', '显示/取片 (2047 回归) + lz4 + 上行发送 (→GW5A)', S_DO, 40, '', '', '', '', '', ''),
 ('1.2.1', '', '2047 显示位流 (M90)', '', '构建 + 加载 + 上屏', S_DO, 60, '自测→系统', 'DR1M90 + 2047 屏', '构建可脱离', '否', 'D1 bin md5 2ac42cbb; D2 PASS', ''),
 ('1.2.1.a', '', '', '单面验证 (D4)', '三色循环 + 数字棋盘格, 低亮度起 (用户定流程)', S_DO, 70, '系统', '2047 屏 (DR1M90-02)', '否', '否', '10-10: Bank11 跳帽改 3.3V 后 s1 ✅、三色循环 ✅; 数字棋盘格已上屏待看', '看棋盘格 → 双面'),
 ('1.2.1.b', '', '', '双面全白 (D5)', 'dual_en, PHASE_B=180, mirror_a/b=0, 无花屏', S_NO, 0, '系统', '2047 屏', '否', '否', '脚本无屏空跑通过', '等屏'),
 ('1.2.1.c', '', '', '3-bit / 偏心几何 (D8/D9)', '棋盘/斜坡与 7020 同图; 转起来体像方向对', S_NO, 0, '系统', '2047 屏 + 电机', '否', '否', 'cmds_3bit_working.txt (V90 09-01)', ''),
 ('1.2.2', '', 'PL lz4 三引擎 (M90)', '', '解压正确 + 并发不扰取片', S_DO, 40, '自测', 'DR1M90', '否', '是', 'D2 自检: NENG=3, 寄存器读写对', 'D10; 可先在板上单独验解压'),
 ('1.2.3', '', '上行发送 uplink_tx (→GW5A)', '', '13 lane 真帧 + 转发时钟, 800M / 1.0G; M90 位流', S_DO, 70, '自测→合测', '合测: DR1M90+ACG525 飞线 (换 0Ω 排阻, J2=1.8V)', '综合/仿真已离线完成', '是', 'uplink_tx/m90: 800M/1.0G bin 出, pclk 余量 46%/26%; 仿真接 GW5A 真接收端逐字节 PASS (dr1v90 c73a5a7)', 'D13/X2 上板'),
 # ===== 1.3 GW5A-25 桥片 =====
 ('1.3', 'GW5A-25 · 桥片固件', '', '', '上行接收 → 分发 → 2260 序列器 → tx 阵列; 两模式', S_DO, 55, '', '', '', '', '', ''),
 ('1.3.1', '', '上行接收 rtl_x8', '', 'IDES8 + 对齐 + deskew', S_OK, 100, '自测', 'ACG525', '仿真可', '是', 'gw5a_loop 4lane@400M 硅上 0 错; bench 13 lane 片内', '800M 长线待短绞线'),
 ('1.3.2', '', '分发 + 帧门控', '', '逐像素正确; 刷新不坏帧 (UL_DROP)', S_OK, 100, '自测', 'ACG525', '仿真可', '是', 'gw5a_bench 1025 万帧 0 坏, 0 retrain', ''),
 ('1.3.3', '', '2260 序列器 + 寄存器表', '', '配置广播一致; 寄存器表 pROM 初始化对', S_OK, 100, '自测', 'ACG525', '仿真可', '是', '签名 1ED7F783/ACEED86D = 仿真', ''),
 ('1.3.4', '', 'tx 阵列 + ODDR 出线', '', 'mini-LVDS 出线逐比特', S_OK, 100, '自测', 'ACG525', '仿真可', '是', '链0 300M 过跳线 19 min 0 错', ''),
 ('1.3.5', '', '验证板配置 (60 扫 1串2)', '', '每面 6 链×2 颗, 18 数据对 + 6 时钟对; 4,800 字槽/链/帧; DCLK 166.67; 4,296 fps', S_DO, 67, '自测', '上样板需 2260 样板', '是', '是', 'build_gw5a.sh 25S "" VALID; README 验证板 profile (dr1v90 f088f31)', '1.3.5.c 上样板 (D15)'),
 ('1.3.5.a', '', '', '参数化 + 仿真对拍', '独立解码器逐像素; 级联语义; 刷新压力档', S_OK, 100, '自测', '无', '是', '是', '5 真帧逐字全对 + VCRC; 每 3 帧刷新也无坏帧; 目标板回归逐行相同', ''),
 ('1.3.5.b', '', '', 'DCLK 166.67 时序收敛', '整桥 166.67 MHz 0 违例', S_OK, 100, '自测', '无', '是', '是', 'dclk 166.871 MHz, LUT 5179, BSRAM 48/56', '⚠ 4,296 fps 余量仅 0.2%, 可把 IDLE_CLK 64→32 换 0.6 µs (需上板确认)'),
 ('1.3.5.c', '', '', '2260 样板出图 (D15)', '已知图案; 1串2 写入', S_NO, 0, '合测', 'GW5A + 2260 样板', '否', '否', '', ''),
 ('1.3.6', '', '压测模式', '', '48 扫帧长 / 603 片 / 假 24 链满载, 误码可见', S_NO, 0, '自测', '无 (仿真) / ACG525', '是', '是', '', 'D17'),
 ('1.3.7', '', '真引脚 (转接板网表)', '', '按逐球引脚分配重排 IO_LOC, 出 bit', S_NO, 0, '自测', '无', '是', '是', 'POV3D_GW5A转LED转接板_逐球引脚分配_2026-09-16', '可先做'),
 ('1.3.12', '', '跨板接收时钟', '', 'PLL 参考 = DR1 转发时钟 (线速/8, 鉴相 ÷2); IODELAY + IDES8 + 扫眼 + 极性', S_OK, 100, '自测→合测', '板上验证并入 D13', '是', '是', '仿真: 800M/1.0G +100 ppm 各 10 万 beat 0 错; ±300 ppm、PLL 相位扫参全过; 反例本地参考链路起不来, 2 ppm 下 7.4 万错 (dr1v90 419a60e)', '上板并入 D13'),
 ('1.3.13', '', 'ACG525 上行接收工程 (X2)', '', '13 对 P6 + 转发时钟 P11-20, 后接出货桥核心 + vid_check; 板上 train → 等 S 行 done → go', S_DO, 70, '合测', 'DR1M90 + ACG525 飞线 (FS03 换 0Ω, J2=1.8V)', '构建已完成', '否', 'BIT OK pclk 127 / dclk 151.6 MHz; README gw5a_xboard/README.md', 'D13 上板'),
 ('1.3.14', '', '目标板帧间空闲 vs 链路看门狗', '', 'DR1 帧间空闲 5.3 µs 正好卡在 lane_align 掉锁看门狗 (5.12 µs) 边上, 降速即每帧重训练; 评估目标板默认开 UL_IDLE_OK', S_NO, 0, '自测', '无', '是', '是', 'agent A 10-10 在验证板 profile 上发现并修 (UL_IDLE_OK)', '评估'),
 ('1.3.8', '', '刷新不丢帧 (X3)', '', '分块刷新, 帧号跳 0', S_NO, 0, '自测', 'ACG525', '是', '是', '现每秒干净丢 1 帧', '待定, 放 M3 后'),
 ('1.3.9', '', 'lane 47 单字错', '', '判定量具问题还是桥问题', S_DO, 70, '自测', 'ACG525', '否', '是', '改量具后 1.3 h 0 错, 长跑中', '读计数'),
 ('1.3.10', '', 'ACK 回读', '', '2260 回包解析', S_BLK, 0, '合测', '2260', '否', '否', 'project_icnd2260_ack_never_received', ''),
 ('1.3.11', '', 'Flash 冷启动', '', '', S_SKIP, 0, '自测', 'ACG525', '否', '是', '用户定跳过 (Flash 内容不明)', ''),
 # ===== 1.4 PC 上位机 =====
 ('1.4', 'PC · 上位机', '', '', '渲染/切片/体素表/推流', S_DO, 40, '', '', '', '', '', ''),
 ('1.4.1', '', '2047 现役管线', '', '7020 上已出体像', S_OK, 100, '系统', '7020 台架', '-', '-', 'project_pov3d_dual_panel_volumetric_ok', ''),
 ('1.4.2', '', '验证板口径推流', '', '160×180, 358 片 @12 rev/s; 体素表 VXT1 (不传全切片), PVS1 flags=VOXTAB|LZ4', S_OK, 100, '自测', '无', '是', '是', 'povstream --profile vb358; 10 帧×2 面×358 片 e2e 对拍 PASS; 最重真实内容 175 KB/帧 = 16.8 Mbps @12 fps (WiFi 159 余 9.4×)', '板端接收/切片 → 1.1.10'),
 # ===== 1.5 系统联调 =====
 ('1.5', '系统联调 (四门)', '', '', '', S_DO, 15, '系统', '', '', '', '', ''),
 ('1.5.1', '', 'M1 2047 双面全白', '', 'D5', S_NO, 0, '系统', 'DR1M90 + 2047 屏', '否', '否', '', '等屏 (焊接 ~2 h)'),
 ('1.5.2', '', 'M2 推流体像', '', 'D10, 10–12 fps 档', S_NO, 0, '系统', '+ 电机 + WiFi', '否', '否', '', ''),
 ('1.5.3', '', 'M3 主控↔桥环回', '', 'D13, ≥3 min 零误码 (800M + 1.0G)', S_NO, 0, '合测', 'DR1M90 + ACG525 + 飞线', '否', '否', '', ''),
 ('1.5.4', '', 'M4 两模式 + 开机 + 远程', '', 'D20', S_NO, 0, '系统', '全套', '否', '否', '', ''),
 ('1.5.6', '', '转接板改版前提 (软件依赖的硬件项)', '', 'UL_CLK 改回 B4_B33 (GCLK/PLL 脚); JTAG 加 2.5↔3.3V 转换 + 下载座; FPC 尾部 64~72 按契约重排', S_NO, 0, '硬件', '转接板下一版', '否', '否', '改板结论 09-18 必改项, 10-10 补注"上行前提 / 远程升级前提"; 规格书 v3 10-10 修订 (UL_CLK 线速/8)', '硬件同事改板时必须落实'),
 ('1.5.5', '', '验证板整机上转子', '', '新灯板 + 转接板 + 转子 + WPT', S_BLK, 0, '系统', '硬件未回', '否', '否', '不在 20 日计划内', ''),
]

HEAD = ['编号', '2级 模块', '3级 功能', '4级 子项', '内容 / 判据', '状态', '完成度%', '验证层级', '依赖硬件',
        '可脱离硬件?', '可独立验证?', '证据 / 记录', '下一步']
FILL = {S_OK: 'C6EFCE', S_DO: 'FFEB9C', S_NO: 'EDEDED', S_BLK: 'FFC7CE', S_SKIP: 'D9D9D9'}
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
ws = wb.active; ws.title = '1级目标'
ws.append(['1级目标', '验证板显示完成 (DR1M90 + 2×GW5A-25 + 2 块 160×180 2260 灯板, 转起来出 3D 体像)', ''])
ws.append([]); ws.append(['参数', '值 / 推演', '出处'])
for r in PARAMS: ws.append(list(r))
for c, w in zip('ABC', (22, 90, 46)): ws.column_dimensions[c].width = w
for row in ws.iter_rows():
    for cell in row: cell.alignment = WRAP
ws['A1'].font = ws['B1'].font = Font(bold=True, size=13)
for cell in ws[3]: cell.font = Font(bold=True)

t = wb.create_sheet('任务分解')
t.append(HEAD)
for r in ROWS:
    t.append(list(r))
widths = (9, 20, 24, 22, 46, 10, 9, 11, 26, 12, 11, 44, 26)
for i, w in enumerate(widths, 1): t.column_dimensions[get_column_letter(i)].width = w
for cell in t[1]:
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='305496'); cell.alignment = WRAP
for row in t.iter_rows(min_row=2):
    lvl = row[0].value.count('.')
    for cell in row:
        cell.alignment = WRAP; cell.border = BOX
        if lvl == 1: cell.font = Font(bold=True)
    st = row[5].value
    if st in FILL: row[5].fill = PatternFill('solid', fgColor=FILL[st])
    if lvl == 1:
        for cell in row:
            if cell.column != 6: cell.fill = PatternFill('solid', fgColor='DDEBF7')
t.freeze_panes = 'B2'
t.auto_filter.ref = t.dimensions

s = wb.create_sheet('汇总')
s.append(['2级模块', '3级项数', '✅完成', '🔶进行中', '⬜未开始', '⛔阻塞/⏸跳过', '可脱离硬件且未完成 (可先跑)'])
for top in [r for r in ROWS if r[0].count('.') == 1]:
    kids = [r for r in ROWS if r[0].startswith(top[0] + '.') and r[0].count('.') == 2]
    cnt = lambda st: sum(1 for k in kids if k[5] == st)
    free = [k[0] + ' ' + k[2] for k in kids if k[5] in (S_NO, S_DO) and k[9] not in ('否', '-')]
    s.append([top[1], len(kids), cnt(S_OK), cnt(S_DO), cnt(S_NO), cnt(S_BLK) + cnt(S_SKIP), '\n'.join(free)])
for c, w in zip('ABCDEFG', (22, 9, 8, 9, 9, 12, 60)): s.column_dimensions[c].width = w
for row in s.iter_rows():
    for cell in row: cell.alignment = WRAP
for cell in s[1]: cell.font = Font(bold=True)

wb.save(OUT)
print('saved', OUT, len(ROWS), 'rows')
